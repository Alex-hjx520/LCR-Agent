"""DOCX 解析器：基于 python-docx，保留段落样式与表格。"""

from __future__ import annotations

import logging
from pathlib import Path

from parsing.base import BaseParser, RawBlock
from schemas.document import SourceType

logger = logging.getLogger(__name__)

#: 会被识别为标题的样式名前缀
_HEADING_STYLE_PREFIXES = ("Heading", "Title", "标题", "Subtitle")


class DocxParser(BaseParser):
    """提取 DOCX 的段落与表格，尽量保留文档结构信息。"""

    source_type = SourceType.DOCX
    suffixes = (".docx",)

    def __init__(self, *, include_tables: bool = True) -> None:
        self.include_tables = include_tables

    def _iter_blocks(self, path: Path) -> list[RawBlock]:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(path))
        blocks: list[RawBlock] = []

        for item in _iter_block_items(document):
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if not text:
                    continue
                style_name = item.style.name if item.style is not None else None
                style = None
                if style_name:
                    style = "heading" if _is_heading_style(style_name) else "body"
                blocks.append((text, None, None, style))
            elif isinstance(item, Table):
                if not self.include_tables:
                    continue
                md = _docx_table_to_markdown(item)
                if md:
                    blocks.append((md, None, None, "table"))

        logger.debug("DOCX 解析完成，共 %d 块", len(blocks))
        return blocks


def _is_heading_style(style_name: str) -> bool:
    return any(style_name.startswith(p) for p in _HEADING_STYLE_PREFIXES)


def _iter_block_items(parent):
    """按文档真实顺序交替产出 Paragraph 与 Table。

    python-docx 的 `document.paragraphs` / `document.tables` 会打乱原始顺序，
    因此这里遍历 body 的 XML 子元素。
    """
    from docx.document import Document as _Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    else:  # _Cell 等容器
        parent_elm = parent._tc  # noqa: SLF001

    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _docx_table_to_markdown(table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        rows.append(cells)
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, *body = rows
    out = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    out.extend("| " + " | ".join(r) + " |" for r in body)
    return "\n".join(out)


__all__ = ["DocxParser"]
